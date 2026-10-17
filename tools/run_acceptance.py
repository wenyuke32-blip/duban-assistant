"""可复现的双轨验收。仅本地文件与显式模拟，无真实邮件/模型网络调用。
运行 python tools/run_acceptance.py --output /隔离验收目录
"""
import argparse, hashlib, json, os, subprocess, sys
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from openpyxl import Workbook,load_workbook
from docx import Document
PROJECT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(PROJECT))
from duban.store import Store, DEPTS
TZ=timezone(timedelta(hours=8))

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--output',required=True);args=parser.parse_args()
    out=Path(args.output).resolve()
    if out.exists() and any(out.iterdir()): raise SystemExit('验收目录必须为空，避免影响现有文件')
    out.mkdir(parents=True,exist_ok=True); fixtures=out/'fixtures';fixtures.mkdir();today=datetime.now(TZ).date()
    trace=[];checks=[]
    def check(category,name,condition,detail=None):
        checks.append({'category':category,'check':name,'pass':bool(condition),'detail':detail})
        if not condition: raise AssertionError(name+': '+str(detail))
    def write(name,obj):
        p=fixtures/name;p.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8');return str(p)
    def cli(workspace,*arguments,sim=False,expected=0):
        cmd=[sys.executable,'-m','duban',*map(str,arguments),'--workspace',str(workspace)]
        env=dict(os.environ)
        if sim:cmd+=['--adapter','tests.fixture_adapter'];env['DUBAN_FIXTURE_DIR']=str(fixtures)
        proc=subprocess.run(cmd,cwd=PROJECT,env=env,capture_output=True,text=True)
        try: payload=json.loads(proc.stdout)
        except Exception: raise AssertionError(proc.stdout+'\n'+proc.stderr)
        trace.append({'args':list(map(str,arguments)),'workspace':workspace.name,'simulation':sim,'exit_code':proc.returncode,'result':payload})
        if proc.returncode!=expected: raise AssertionError(payload)
        return payload
    rows=[
      ['测试渠道走访','市场部',(today+timedelta(days=7)).isoformat(),'未开始',''],
      ['测试企业专线交付','政企部',(today-timedelta(days=1)).isoformat(),'未开始',''],
      ['测试园区扩容','网络部',(today+timedelta(days=3)).isoformat(),'办理中',''],
      ['测试流程材料','综合部',(today+timedelta(days=1)).isoformat(),'已完成','材料已归档'],
      ['测试客户走访','政企部',(today+timedelta(days=2)).isoformat(),'未开始',''],
      ['测试取消事项','综合部',(today+timedelta(days=1)).isoformat(),'已取消','已取消']]
    wb=Workbook();ws=wb.active;ws.append(['事项名称','责任部门','截止日期','状态','进展'])
    for row in rows:ws.append(row)
    xlsx=fixtures/'input.xlsx';wb.save(xlsx);wb.close()
    def seed(workspace):
        cli(workspace,'init');result=cli(workspace,'import','--file',xlsx)
        assert result['结果']['新增']==6
        for i,d in enumerate(DEPTS):
            path=write('contact.json',{'部门名称':d,'主管邮箱':f'd{i}@example.invalid','管理员邮箱':'admin@example.invalid'})
            cli(workspace,'contacts','--file',path)
        return cli(workspace,'query')['结果']['tasks']
    def inspect_reports(workspace,category):
        reports=workspace/'reports';files=list(reports.iterdir());check(category,'10份报告文件真实存在',len(files)==10)
        raw=Store(workspace).read();tasks=raw['active']+raw['closed']
        for scope in DEPTS+['管理员综合']:
            own=tasks if scope=='管理员综合' else [t for t in tasks if t['主责部门']==scope]
            book=load_workbook(reports/(scope+'督办报告.xlsx'));metrics={r[0]:r[1] for r in list(book['可视化看板'].values)[1:]}
            check(category,scope+'数量与真实台账独立核对',metrics['任务总数']==len(own))
            check(category,scope+'四个Excel原生图表',len(book['可视化看板']._charts)==4)
            status={r[0]:r[1] for r in list(book['状态分布'].values)[1:]};expected=Counter(t['办理状态'] for t in own)
            check(category,scope+'图表源数据正确',all(status[k]==expected[k] for k in status))
            book.close();doc=Document(reports/(scope+'督办报告.docx'));text='\n'.join(p.text for p in doc.paragraphs)
            check(category,scope+'Word图表及正文',len(doc.inline_shapes)>=1 and f'共有任务{len(own)}项' in text)
            for t in tasks:
                if scope!='管理员综合' and t['主责部门']!=scope and scope not in t.get('协办部门','').split('、'):
                    check(category,scope+'不含其他部门专属任务:'+t['任务名称'],t['任务名称'] not in text)
    local=out/'local-workspace';tasks=seed(local);ids={t['任务名称']:t['内部任务唯一ID'] for t in tasks}
    check('1','跨独立进程初始化不覆盖任务',len(cli(local,'init')['结果'])>0 and {t['内部任务唯一ID'] for t in cli(local,'query')['结果']['tasks']}==set(ids.values()))
    progress=write('progress.json',{'办理状态':'受阻','最新办理进展':'两处资源不足','预计完成日期':(today+timedelta(days=10)).isoformat()})
    cli(local,'update','--task-id',ids['测试企业专线交付'],'--file',progress)
    result=cli(local,'update','--task-id',ids['测试企业专线交付'],'--file',progress)
    check('1','重复相同更新不产生业务更新',result['执行摘要']['本次更新任务数']==0)
    before=next(t for t in tasks if t['任务名称']=='测试渠道走访')['原定截止日期']
    cli(local,'close','--task-id',ids['测试渠道走访']);cli(local,'reopen','--task-id',ids['测试渠道走访'])
    reopened=next(t for t in cli(local,'query')['结果']['tasks'] if t['内部任务唯一ID']==ids['测试渠道走访'])
    check('1','归档重开保留ID和原截止',reopened['重新开启次数']==1 and reopened['原定截止日期']==before and reopened['实际完成日期']=='')
    cli(local,'reports');inspect_reports(local,'1');first_target=(local/'reports').resolve()
    hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in first_target.iterdir()}
    check('1','五封报告正确入队',len(cli(local,'send-reports')['结果']['入队'])==5)
    check('1','相同报告再次生成不重复入队',len(cli(local,'reports')['结果'])==10 and cli(local,'send-reports')['结果']['入队']==[])
    check('1','未配置邮箱时不会报告发送成功',cli(local,'dispatch')['结果']['status']=='blocked')
    check('1','重复报告保持同一版本及全部字节不变',(local/'reports').resolve()==first_target and hashes=={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in first_target.iterdir()})
    # 子进程真实退出，发生于报告已切换但台账尚未提交之间。
    code='from duban.store import Store; from duban import core; from duban.reports import generate; import os,sys\ns=Store(sys.argv[1])\nwith s.transaction():\n core.update(s,sys.argv[2],{"最新办理进展":"进程中断前未提交内容"})\n generate(s)\n os._exit(19)'
    dead=subprocess.run([sys.executable,'-c',code,str(local),ids['测试企业专线交付']],cwd=PROJECT,capture_output=True,text=True)
    check('1','报告发布窗口进程中断已触发',dead.returncode==19 and (local/'.report-publication.json').exists())
    cli(local,'query')
    check('1','重启后报告及台账回到同一有效版本',(local/'reports').resolve()==first_target and not (local/'.report-publication.json').exists() and hashes=={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in first_target.iterdir()})
    # 两个真实命令行进程同时新增，验证文件锁避免丢失更新。
    jobs=[]
    for i in range(2):
        p=write(f'concurrent{i}.json',{'任务名称':f'测试并发事项{i}','主责部门':'综合部','原定截止日期':(today+timedelta(days=30)).isoformat()})
        jobs.append(subprocess.Popen([sys.executable,'-m','duban','add','--workspace',str(local),'--file',p],cwd=PROJECT,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True))
    codes=[job.communicate()+(job.returncode,) for job in jobs]
    check('1','并发新增无丢失更新',all(x[2]==0 for x in codes) and len(cli(local,'query')['结果']['tasks'])==8)
    cli(local,'reports');local.joinpath('督办管理总台账.xlsx').write_bytes(b'CORRUPTED-BY-ACCEPTANCE')
    cli(local,'query',expected=1)
    check('1','损坏文件不会被重新初始化',local.joinpath('督办管理总台账.xlsx').read_bytes()==b'CORRUPTED-BY-ACCEPTANCE')
    recovery=cli(local,'recover')['结果'];check('1','显式恢复备份并保留损坏副本',recovery['恢复任务数']==8 and Path(recovery['原文件保留位置']).read_bytes()==b'CORRUPTED-BY-ACCEPTANCE')
    check('1','恢复后仍可读取八项任务',len(cli(local,'query')['结果']['tasks'])==8)
    # 显式模拟宿主：模型结果预先写定，邮件为本地JSON，分发为本地文件复制。
    sim=out/'simulation-workspace';tasks=seed(sim);ids={t['任务名称']:t['内部任务唯一ID'] for t in tasks}
    write('allowed-recipients.json',[f'd{i}@example.invalid' for i in range(4)]+['admin@example.invalid'])
    config=write('issuer.json',{'正式下达邮箱':['d1@example.invalid']});cli(sim,'configure','--file',config)
    deadline=(today+timedelta(days=30)).isoformat();quote=f'正式督办：政企部测试企业信息摸排，{deadline}完成。'
    body='测试企业专线交付两处资源不足；测试客户走访已完成。'+quote
    doc=Document();doc.add_paragraph('测试园区扩容资源已到位，完成60%');doc.save(fixtures/'feedback.docx')
    sent=datetime.now(TZ).isoformat(timespec='seconds')
    messages=[{'id':'SIM-1','sender':'d1@example.invalid','sent_at':sent,'body':body},
      {'id':'SIM-2','sender':'d2@example.invalid','sent_at':sent,'body':'','attachments':[{'name':'feedback.docx','local_path':str(fixtures/'feedback.docx')}]},
      {'id':'SIM-3','sender':'d1@example.invalid','sent_at':sent,'body':'测试任务尚未完成，请保留待确认。'}]
    def item(name,evidence,fields,source='body'):
        return {'action':'update','task_id':ids[name],'candidate_ids':[ids[name]],'evidence':[evidence],'anchors':[name],'fields':fields,'source_id':source}
    plans={'SIM-1':[item('测试企业专线交付','测试企业专线交付两处资源不足',{'办理状态':'受阻','最新办理进展':'两处资源不足'}),item('测试客户走访','测试客户走访已完成',{'办理状态':'已完成'}),{'action':'create','issuance_quote':quote,'evidence':[quote],'fields':{'任务名称':'测试企业信息摸排','主责部门':'政企部','原定截止日期':deadline}}],
      'SIM-2':[item('测试园区扩容','测试园区扩容资源已到位，完成60%',{'办理状态':'办理中','完成百分比':60,'最新办理进展':'资源已到位'},'attachment:0')],
      'SIM-3':[{'action':'update','task_id':ids['测试企业专线交付'],'candidate_ids':[ids['测试企业专线交付'],ids['测试客户走访']],'evidence':['测试任务尚未完成'],'anchors':['测试'],'fields':{'办理状态':'受阻'}}]}
    write('messages.json',messages);write('plans.json',plans)
    result=cli(sim,'full',sim=True)
    check('2','模拟两页三封邮件全部拉取',result['结果']['邮箱同步']['读取邮件数量']==3 and result['结果']['邮箱同步']['页数']==2)
    check('2','模拟邮件新增更新办结及歧义隔离',result['执行摘要']['本次新增任务数']==1 and result['执行摘要']['本次办结任务数']==1 and result['执行摘要']['本次新增待确认数']==1)
    current=cli(sim,'query')['结果']['tasks'];network=next(t for t in current if t['任务名称']=='测试园区扩容')
    check('2','本地Word附件真实解析后应用模拟计划',network['完成百分比']==60)
    check('2','固定fixture研判结果有原文依据并入台账',any(t['当前AI风险等级']=='高' and t['AI分析依据']=='两处资源不足' for t in current))
    receipts=json.loads((fixtures/'receipts.json').read_text());report_receipts=[r for r in receipts.values() if r['attachments']]
    check('2','五封模拟报告每封两附件且不串部门',len(report_receipts)==5 and all(len(r['attachments'])==2 for r in report_receipts))
    for i,d in enumerate(DEPTS):
        r=next(r for r in report_receipts if r['recipient']==f'd{i}@example.invalid')
        check('2',d+'模拟分发隔离',set(r['attachments'])=={d+'督办报告.xlsx',d+'督办报告.docx'})
    check('2','管理员综合模拟分发隔离',set(next(r for r in report_receipts if r['recipient']=='admin@example.invalid')['attachments'])=={'管理员综合督办报告.xlsx','管理员综合督办报告.docx'})
    check('2','模拟发送不计入真实发送数',result['执行摘要']['本次发送成功数量']==0 and result['执行摘要']['本次模拟发送数量']==len(receipts))
    before_receipts=len(receipts);before_reports=(sim/'reports').resolve();again=cli(sim,'full',sim=True)
    check('2','模拟完整流程重跑不重复投递或重建报告',len(json.loads((fixtures/'receipts.json').read_text()))==before_receipts and again['执行摘要']['本次模拟发送数量']==0 and (sim/'reports').resolve()==before_reports)
    admin=write('admin-plan.json',[{'action':'create','fields':{'任务名称':'测试模拟自然语言新增','主责部门':'网络部','原定截止日期':deadline}}])
    nl=cli(sim,'nl','--text','模拟自然语言输入，实际计划由fixture固定提供',sim=True)
    check('2','模拟自然语言入口只验证计划执行',nl['执行摘要']['本次新增任务数']==1)
    summary={'time':datetime.now(TZ).isoformat(),'passed':sum(c['pass'] for c in checks),'failed':sum(not c['pass'] for c in checks),
      '真实邮箱发送':0,'真实模型调用':0,'checks':checks,'limitations':['MobileWork宿主执行权限/重启持久化未验证','cmcc-mail真实工具绑定未实现','真实邮件语义匹配和AI原因质量未验证','真实附件下载/发送及收件箱投递未验证','正式无人值守授权未核实']}
    (out/'checks.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    (out/'cli-trace.json').write_text(json.dumps(trace,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in summary.items() if k!='checks'},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
