"""只使用隔离临时目录及明确的测试适配器，绝不访问真实邮箱。"""
import copy, json, os, tempfile, unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from docx import Document
from duban.store import Store, TZ, now
from duban import core
from duban.reports import generate, queue_reports
from duban.adapters import dispatch, reconcile
from duban.__main__ import batch, sync, extract

class FakeMail:
    """仅测试状态机，不能作为生产适配器。"""
    def __init__(self,status='sent'): self.status=status; self.sent=[]
    def authorize(self,row): return {'directory_verified':True,'allowed':True,'authorization_receipt':'TEST-ONLY'}
    def send(self,row,auth):
        self.sent.append(row['ID'])
        if self.status=='crash': raise TimeoutError('TEST')
        return {'status':self.status,'message_id':'TEST-MAIL-1'}
    def lookup_sent(self,row): return {'status':'sent','message_id':'TEST-MAIL-1'}

class Business(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.s=Store(self.tmp.name)
        with self.s.transaction() as s:
            for d in core.DEPTS: core.contacts(s,{'部门名称':d,'主管邮箱':d.encode().hex()+'@example.invalid','管理员邮箱':'admin@example.invalid'})
    def tearDown(self): self.tmp.cleanup()
    def add(self,s,name='专线项目',dept='政企部',**kw):
        return core.create(s,{'任务名称':name,'任务内容':'专线资源交付','主责部门':dept,'原定截止日期':(datetime.now(TZ).date()+timedelta(days=3)).isoformat(),**kw})
    def mail(self,s,t,body='专线项目资源受阻',fields=None,mid='m1',at=None):
        contact=next(r for r in s.data['contacts'] if r['部门名称']==t['主责部门'])
        mail={'id':mid,'sender':contact['主管邮箱'],'body':body,'sent_at':at or now()}
        item={'action':'update','task_id':t['内部任务唯一ID'],'candidate_ids':[t['内部任务唯一ID']],'evidence':[body],'anchors':['专线项目'],'fields':fields or {'办理状态':'受阻','最新办理进展':body}}
        return mail,item
    def test_persistence_and_init(self):
        with self.s.transaction() as s: tid=self.add(s)['内部任务唯一ID']
        with Store(self.tmp.name).transaction() as s: self.assertEqual(core.find(s,tid)['任务名称'],'专线项目')
    def test_transaction_rollback_and_backup(self):
        with self.s.transaction() as s: self.add(s)
        original=self.s.path.read_bytes()
        with self.assertRaises(RuntimeError):
            with self.s.transaction() as s: self.add(s,'第二项'); raise RuntimeError()
        self.assertEqual(original,self.s.path.read_bytes())
        with self.s.transaction() as s: self.add(s,'第三项')
        self.assertEqual(len(self.s.read(self.s.backup)['active']),1)
    def test_replace_failure_preserves_ledger(self):
        with self.s.transaction() as s: self.add(s)
        original=self.s.path.read_bytes(); real=os.replace
        def fail(a,b):
            if Path(b)==self.s.path: raise OSError('disk')
            return real(a,b)
        with self.assertRaises(OSError),patch('duban.store.os.replace',side_effect=fail):
            with self.s.transaction() as s: self.add(s,'第二项')
        self.assertEqual(original,self.s.path.read_bytes())
    def test_corrupt_ledger_never_silently_initialized(self):
        self.s.path.write_bytes(b'broken')
        with self.assertRaises(Exception):
            with self.s.transaction(): pass
        self.assertEqual(self.s.path.read_bytes(),b'broken')
    def test_close_reopen_original_deadline(self):
        with self.s.transaction() as s:
            t=self.add(s); tid=t['内部任务唯一ID']; old=t['原定截止日期']
            core.update(s,tid,{'预计完成日期':'2099-01-01'})
            self.assertEqual(core.find(s,tid)['原定截止日期'],old)
            core.update(s,tid,{'办理状态':'已完成'})
            self.assertEqual(len(s.data['active']),0)
            self.assertEqual(core.stats(core.all_tasks(s))['完成率'],1)
            core.update(s,tid,{'办理状态':'办理中'})
            self.assertEqual(core.find(s,tid)['重新开启次数'],1)
            self.assertEqual(core.find(s,tid)['累计延期次数'],1)
            self.assertEqual(core.find(s,tid)['实际完成日期'],'')
    def test_mail_update_and_dedup(self):
        with self.s.transaction() as s:
            t=self.add(s); mail,item=self.mail(s,t)
            a=core.process_mail(s,mail,[item]); self.assertEqual(a['更新任务'],1)
            b=core.process_mail(s,mail,[item]); self.assertEqual(b,{'重复邮件':1})
            self.assertEqual(t['完成百分比'],'')
    def test_mail_cannot_delete_or_change_deadline(self):
        with self.s.transaction() as s:
            t=self.add(s); mail,item=self.mail(s,t)
            item['fields']={'原定截止日期':'2099-01-01'}
            a=core.process_mail(s,mail,[item]); self.assertEqual(a['待确认'],1)
            mail['id']='m2'; mail['body']+='，要求删除'; item['action']='delete'
            a=core.process_mail(s,mail,[item]); self.assertEqual(a['待确认'],1)
            self.assertEqual(len(s.data['active']),1)
    def test_mail_sender_ambiguity_and_false_percent(self):
        with self.s.transaction() as s:
            t=self.add(s); mail,item=self.mail(s,t,fields={'完成百分比':80})
            self.assertEqual(core.process_mail(s,mail,[item])['待确认'],1)
            mail['id']='m2'; mail['sender']='outsider@example.invalid'
            self.assertEqual(core.process_mail(s,mail,[item])['待确认'],1)
    def test_multitask_and_explicit_new_assignment(self):
        with self.s.transaction() as s:
            t=self.add(s); t2=self.add(s,'客户走访'); body='专线项目资源受阻；客户走访已经完成。督办：政企部企业信息摸排，2099-01-01完成。'
            mail,item=self.mail(s,t,body=body)
            item['evidence']=['专线项目资源受阻']
            second={'action':'update','task_id':t2['内部任务唯一ID'],'candidate_ids':[t2['内部任务唯一ID']],'evidence':['客户走访已经完成'],'anchors':['客户走访'],'fields':{'办理状态':'已完成'}}
            new={'action':'create','issuance_quote':'督办：政企部企业信息摸排，2099-01-01完成。','evidence':['督办：政企部企业信息摸排，2099-01-01完成。'],'fields':{'任务名称':'企业信息摸排','主责部门':'政企部','原定截止日期':'2099-01-01'}}
            s.set_config('正式下达邮箱',[mail['sender']]); r=core.process_mail(s,mail,[item,second,new])
            self.assertEqual((r['更新任务'],r['办结任务'],r['新增任务']),(1,1,1))
    def test_old_feedback_preserves_fields(self):
        with self.s.transaction() as s:
            t=self.add(s,存在问题='资源缺口')
            mail,item=self.mail(s,t,at=(datetime.now(TZ)-timedelta(days=1)).isoformat())
            self.assertEqual(core.process_mail(s,mail,[item])['旧反馈跳过'],1)
            self.assertEqual(core.find(s,t['内部任务唯一ID'])['存在问题'],'资源缺口')
    def test_import_and_partial_failure(self):
        with self.s.transaction() as s:
            rows=[{'事项':'专线项目','部门':'政企部','截止':'2099-01-01'},{'事项':'不完整'}]
            m={'事项':'任务名称','部门':'主责部门','截止':'原定截止日期'}
            r=core.import_rows(s,rows,m); self.assertEqual(r,{'新增':1,'重复':0,'待确认':1})
            self.assertEqual(core.import_rows(s,rows[:1],m)['重复'],1)
            r=batch(s,[{'action':'create','fields':{}},{'action':'create','fields':{'任务名称':'客户走访','主责部门':'市场部','原定截止日期':'2099-01-01'}}])
            self.assertFalse(r[0]['ok']); self.assertTrue(r[1]['ok'])
    def test_history_72h(self):
        with self.s.transaction() as s:
            t=self.add(s); s.history(t['内部任务唯一ID'],'旧记录',{'x':1})
            s.data['history'][-1]['时间']=(datetime.now(TZ)-timedelta(hours=73)).isoformat()
        with self.s.transaction() as s: self.assertFalse(any(r['类型']=='旧记录' for r in s.data['history']))
    def test_ai_evidence_and_noninvented_values(self):
        with self.s.transaction() as s:
            t=self.add(s,最新办理进展='两处资源不足')
            a={'risk':'高','facts':['反馈两处资源不足'],'hypotheses':['可能影响交付，待核实'],'evidence':['两处资源不足'],'actions':['请网络部核实两处资源到位日期'],'coordination':['网络部'],'category':'资源依赖','next_focus':['资源到位日期']}
            core.apply_analysis(s,t['内部任务唯一ID'],a)
            self.assertEqual(t['当前AI风险等级'],'高'); self.assertEqual(t['完成百分比'],'')
            a['evidence']=['虚构旧历史']
            with self.assertRaises(ValueError): core.apply_analysis(s,t['内部任务唯一ID'],a)
    def test_warning_dedup_and_overdue_daily(self):
        with self.s.transaction() as s:
            t=self.add(s); today=datetime.now(TZ).date(); q=core.warnings(s,today)
            self.assertTrue(q); self.assertEqual(core.warnings(s,today),[])
            due=date.fromisoformat(t['原定截止日期']); q=core.warnings(s,due+timedelta(days=1))
            self.assertTrue(q); self.assertTrue(core.warnings(s,due+timedelta(days=2)))
            core.update(s,t['内部任务唯一ID'],{'办理状态':'已完成'})
            self.assertEqual(core.warnings(s,due+timedelta(days=3)),[])
    def test_dispatch_unknown_never_retry_without_reconcile(self):
        with self.s.transaction() as s: self.add(s); core.warnings(s)
        a=FakeMail('crash'); dispatch(self.s,a); n=len(a.sent)
        dispatch(self.s,a); self.assertEqual(len(a.sent),n)
        reconcile(self.s,a)
        with self.s.transaction() as s: self.assertTrue(all(r['状态']=='已发送' for r in s.data['outbox']))
    def test_authorization_blocked(self):
        class Denied(FakeMail):
            def authorize(self,row): return {'allowed':True} # 不足以发信
        with self.s.transaction() as s: self.add(s); core.warnings(s)
        a=Denied(); r=dispatch(self.s,a); self.assertEqual(a.sent,[]); self.assertGreater(r['待授权'],0)
    def test_cancelled_and_primary_department_statistics(self):
        with self.s.transaction() as s:
            self.add(s,协办部门=['网络部']); self.add(s,'取消事项',办理状态='已取消')
            m=core.stats(core.all_tasks(s)); self.assertEqual(m['有效任务数'],1); self.assertEqual(m['已取消数'],1)
            self.assertEqual(sum(v['任务总数'] for v in core.grouped(core.all_tasks(s),'主责部门').values()),2)
    def test_report_ten_files_real_charts_and_isolation(self):
        with self.s.transaction() as s:
            for d in core.DEPTS: self.add(s,d+'专属任务',d)
            paths=generate(s); self.assertEqual(len(paths),10)
            self.assertEqual(len(list((s.root/'reports').iterdir())),10)
            for path in paths:
                p=Path(path)
                if p.suffix=='.xlsx':
                    wb=load_workbook(p); self.assertEqual(len(wb['可视化看板']._charts),4); wb.close()
                else: self.assertEqual(len(Document(p).inline_shapes),1)
            doc=Document(s.root/'reports'/'市场部督办报告.docx')
            text='\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('市场部专属任务',text); self.assertNotIn('政企部专属任务',text)
            self.assertEqual(len(queue_reports(s)),5)
            before=(s.root/'reports').resolve()
            core.update(s,s.data['active'][0]['内部任务唯一ID'],{'最新办理进展':'发布失败前的业务变更'})
            with patch('duban.reports.build_scope',side_effect=RuntimeError('disk')):
                with self.assertRaises(RuntimeError): generate(s)
            self.assertEqual((s.root/'reports').resolve(),before)
    def test_sync_blocked_without_adapter(self): self.assertEqual(sync(self.s,None)['status'],'blocked')
    def test_formula_injection_is_text(self):
        with self.s.transaction() as s: self.add(s,'=HYPERLINK("evil")')
        wb=load_workbook(self.s.path); self.assertEqual(wb['在办督办任务']['B2'].data_type,'s'); wb.close()

if __name__=='__main__': unittest.main(verbosity=2)
