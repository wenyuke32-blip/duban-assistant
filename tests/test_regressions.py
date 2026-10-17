"""审计新增回归：本地真实数据路径/故障注入；不证明模型或邮箱可用。"""
import copy, json, os, tempfile, unittest
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from duban.store import Store, TZ, now
from duban import core
from duban.reports import generate, queue_reports
from duban.adapters import dispatch,reconcile
from duban.__main__ import sync

class Regression(unittest.TestCase):
    def setUp(self): self.tmp=tempfile.TemporaryDirectory(); self.s=Store(self.tmp.name)
    def tearDown(self): self.tmp.cleanup()
    def task(self,s): return core.create(s,{'任务名称':'园区专线','任务内容':'园区专线资源交付','主责部门':'政企部','原定截止日期':'2099-01-01','最新办理进展':'两处资源不足'})
    def contacts(self,s):
        for i,d in enumerate(core.DEPTS): core.contacts(s,{'部门名称':d,'主管邮箱':f'd{i}@example.invalid','管理员邮箱':'admin@example.invalid'})
    def analysis(self): return {'risk':'高','facts':['反馈两处资源不足'],'hypotheses':['资源依赖可能影响交付，待核实'],'evidence':['两处资源不足'],'actions':['网络部核对两处资源到位日期'],'coordination':['网络部'],'category':'资源依赖','next_focus':['到位日期']}
    def test_noop_update_preserves_history_and_queue(self):
        with self.s.transaction() as s:
            t=self.task(s); core.enqueue(s,'a@example.invalid','提醒','正文','k',t['内部任务唯一ID'])
            before=copy.deepcopy(s.data)
            self.assertEqual(core.update(s,t['内部任务唯一ID'],{'最新办理进展':'两处资源不足'}),'无变更')
            self.assertEqual(s.data,before)
    def test_empty_key_fields_rejected_before_corruption(self):
        with self.s.transaction() as s:
            t=self.task(s)
            for field in ['任务名称','主责部门','原定截止日期']:
                for v in ['',None]:
                    with self.assertRaises(ValueError): core.update(s,t['内部任务唯一ID'],{field:v})
    def test_invalid_text_fails_one_batch_item_only(self):
        from duban.__main__ import batch
        with self.s.transaction() as s:
            results=batch(s,[{'action':'create','fields':{'任务名称':'坏\x00数据','主责部门':'综合部','原定截止日期':'2099-01-01'}},{'action':'create','fields':{'任务名称':'有效事项','主责部门':'综合部','原定截止日期':'2099-01-01'}}])
            self.assertEqual([r['ok'] for r in results],[False,True])
        self.assertEqual(len(self.s.read()['active']),1)
    def test_blank_ai_evidence_is_not_evidence(self):
        with self.s.transaction() as s:
            t=self.task(s)
            for e in ['', '   ', '\n']:
                a=self.analysis();a['evidence']=[e]
                with self.assertRaises(ValueError): core.apply_analysis(s,t['内部任务唯一ID'],a)
    def test_update_invalidates_same_second_ai_result(self):
        with self.s.transaction() as s:
            t=self.task(s); core.apply_analysis(s,t['内部任务唯一ID'],self.analysis())
            self.assertIsNotNone(core.current_analysis(s,t))
            core.update(s,t['内部任务唯一ID'],{'最新办理进展':'资源已到位'})
            self.assertIsNone(core.current_analysis(s,t)); self.assertEqual(t['当前AI风险等级'],'未研判')
            self.assertEqual(core.scan(s)['需宿主AI研判'],[t['内部任务唯一ID']])
            self.assertEqual(core.warnings(s),[])
    def test_repeated_pending_is_idempotent(self):
        with self.s.transaction() as s:
            a=s.pending('接口失败',{'类型':'TimeoutError'}); b=s.pending('接口失败',{'类型':'TimeoutError'})
            self.assertEqual(a,b);self.assertEqual(len(s.data['pending']),1)
    def test_percent_80_is_not_evidenced_by_180(self):
        with self.s.transaction() as s:
            self.contacts(s);t=self.task(s);body='园区专线完成180%'
            mail={'id':'percent','sender':'d1@example.invalid','sent_at':now(),'body':body}
            item={'action':'update','task_id':t['内部任务唯一ID'],'candidate_ids':[t['内部任务唯一ID']],'evidence':[body],'anchors':['园区专线'],'fields':{'完成百分比':80}}
            self.assertEqual(core.process_mail(s,mail,[item])['待确认'],1)
            self.assertEqual(core.find(s,t['内部任务唯一ID'])['完成百分比'],'')
    def test_mail_cannot_clear_existing_problem(self):
        with self.s.transaction() as s:
            self.contacts(s);t=self.task(s);body='园区专线最新反馈'
            mail={'id':'clear','sender':'d1@example.invalid','sent_at':now(),'body':body}
            item={'action':'update','task_id':t['内部任务唯一ID'],'candidate_ids':[t['内部任务唯一ID']],'evidence':[body],'anchors':['园区专线'],'fields':{'存在问题':''}}
            self.assertEqual(core.process_mail(s,mail,[item])['待确认'],1)
    def test_duplicate_content_with_new_mail_id_not_applied_twice(self):
        with self.s.transaction() as s:
            self.contacts(s); t=self.task(s); body='园区专线受阻'
            mail={'id':'m1','sender':'d1@example.invalid','sent_at':now(),'body':body}
            item={'action':'update','task_id':t['内部任务唯一ID'],'candidate_ids':[t['内部任务唯一ID']],'evidence':[body],'anchors':['园区专线'],'fields':{'办理状态':'受阻'}}
            core.process_mail(s,mail,[item]);mail['id']='m2'
            self.assertEqual(core.process_mail(s,mail,[item]),{'重复邮件':1})
            self.assertEqual(len(s.data['mail']),2)
    def test_noop_mail_advances_watermark_prevents_older_override(self):
        with self.s.transaction() as s:
            t=self.task(s); at=datetime.now(TZ)+timedelta(seconds=4)
            self.assertEqual(core.update(s,t['内部任务唯一ID'],{'最新办理进展':t['最新办理进展']},'邮件:m1',at.isoformat(),False),'无变更')
            outcome=core.update(s,t['内部任务唯一ID'],{'最新办理进展':'较旧资料'},'邮件:m2',(at-timedelta(seconds=2)).isoformat(),False)
            self.assertEqual(outcome,'旧反馈已跳过')
    def test_corrupt_primary_explicit_restore_preserves_corrupt_copy(self):
        with self.s.transaction() as s: self.task(s)
        with self.s.transaction(): pass
        self.s.path.write_bytes(b'CORRUPTED')
        result=self.s.recover(); self.assertEqual(result['恢复任务数'],1)
        self.assertEqual(Path(result['原文件保留位置']).read_bytes(),b'CORRUPTED')
        self.assertEqual(len(self.s.read()['active']),1)
    def test_unrecognised_sheet_is_not_silently_destroyed(self):
        with self.s.transaction() as s:self.task(s)
        wb=load_workbook(self.s.path);wb.create_sheet('人工业务资料');wb.save(self.s.path);wb.close(); before=self.s.path.read_bytes()
        with self.assertRaises(ValueError):
            with self.s.transaction(): pass
        self.assertEqual(self.s.path.read_bytes(),before)
    def test_repeated_report_and_queue_reuse_exact_files(self):
        with self.s.transaction() as s:self.contacts(s);self.task(s);generate(s);first=queue_reports(s)
        old=(self.s.root/'reports').resolve(); contents={p.name:p.read_bytes() for p in old.iterdir()}
        with self.s.transaction() as s:
            generate(s);self.assertEqual(queue_reports(s),[])
        self.assertEqual((self.s.root/'reports').resolve(),old)
        self.assertEqual(contents,{p.name:p.read_bytes() for p in old.iterdir()});self.assertEqual(len(first),5)
    def test_old_reports_blocked_after_business_update(self):
        with self.s.transaction() as s:self.contacts(s);t=self.task(s);generate(s)
        with self.s.transaction() as s:
            core.update(s,t['内部任务唯一ID'],{'最新办理进展':'资源已到位'})
            with self.assertRaises(ValueError):queue_reports(s)
    def test_report_publish_and_ledger_commit_rollback_together(self):
        with self.s.transaction() as s:self.contacts(s);t=self.task(s);generate(s)
        old=(self.s.root/'reports').resolve(); original=self.s.path.read_bytes(); real=os.replace
        def fail(a,b):
            if Path(b)==self.s.path: raise OSError('injected disk failure')
            return real(a,b)
        with self.assertRaises(OSError),patch('duban.store.os.replace',side_effect=fail):
            with self.s.transaction() as s:
                core.update(s,t['内部任务唯一ID'],{'最新办理进展':'新变化'});generate(s)
        self.assertEqual(self.s.path.read_bytes(),original)
        self.assertEqual((self.s.root/'reports').resolve(),old)
        self.assertFalse((self.s.root/'.report-publication.json').exists())
    def test_malformed_send_response_marks_unknown_and_continues(self):
        class Host:
            EXECUTION_MODE='simulation'
            def authorize(self,row): return {'directory_verified':True,'allowed':True,'authorization_receipt':'SIMULATION'}
            def send(self,row,auth): return None if row['去重键']=='first' else {'status':'sent','message_id':'SIM-2'}
        with self.s.transaction() as s:
            core.enqueue(s,'a@example.invalid','测试','正文','first');core.enqueue(s,'b@example.invalid','测试','正文','second')
        result=dispatch(self.s,Host());self.assertEqual(result['结果未知'],1);self.assertEqual(result['发送成功'],1)
        with self.s.transaction() as s:self.assertEqual([r['状态'] for r in s.data['outbox']],['结果未知','已发送'])
    def test_fetch_exception_is_blocking_not_success(self):
        class Host:
            def fetch_mail(self,cursor): raise TimeoutError()
        self.assertEqual(sync(self.s,Host())['status'],'incomplete')
    def test_nonadvancing_pagination_is_blocked(self):
        class Host:
            def fetch_mail(self,cursor):return {'messages':[],'next_cursor':cursor,'has_more':True}
        self.assertEqual(sync(self.s,Host())['status'],'incomplete')

    def test_equal_time_conflicting_mail_is_pending(self):
        with self.s.transaction() as s:
            self.contacts(s); t=self.task(s); at=now()
            def apply(mid,body,status):
                return core.process_mail(s,{'id':mid,'sender':'d1@example.invalid','sent_at':at,'body':body},[{'action':'update','task_id':t['内部任务唯一ID'],'candidate_ids':[t['内部任务唯一ID']],'evidence':[body],'anchors':['园区专线'],'fields':{'办理状态':status}}])
            self.assertEqual(apply('a','园区专线受阻','受阻')['更新任务'],1)
            self.assertEqual(apply('b','园区专线已完成','已完成')['待确认'],1)
            self.assertEqual(core.find(s,t['内部任务唯一ID'])['办理状态'],'受阻')
    def test_all_warning_calendar_nodes(self):
        with self.s.transaction() as s:
            self.contacts(s);t=self.task(s);due=datetime.fromisoformat(t['原定截止日期']).date()
            for days,expected in [(7,'关注提醒'),(3,'临期预警'),(1,'紧急预警'),(0,None),(-1,'逾期督办')]:
                ids=core.warnings(s,due-timedelta(days=days))
                levels={r['等级'] for r in s.data['outbox'] if r['ID'] in ids}
                self.assertEqual(levels,{expected} if expected else set())
    def test_shared_anchor_is_not_unique_evidence(self):
        with self.s.transaction() as s:
            self.contacts(s);t=self.task(s)
            core.create(s,{'任务名称':'园区专线二期','主责部门':'政企部','原定截止日期':'2099-01-01'})
            body='园区专线受阻'; item={'action':'update','task_id':t['内部任务唯一ID'],'candidate_ids':[t['内部任务唯一ID']],'evidence':[body],'anchors':['园区专线'],'fields':{'办理状态':'受阻'}}
            self.assertEqual(core.process_mail(s,{'id':'ambiguous','sender':'d1@example.invalid','sent_at':now(),'body':body},[item])['待确认'],1)
            self.assertEqual(core.find(s,t['内部任务唯一ID'])['办理状态'],'未开始')
    def test_definitive_failure_can_retry_once_and_then_dedup(self):
        class Host:
            EXECUTION_MODE='simulation'
            calls=0
            def authorize(self,row):return {'directory_verified':True,'allowed':True,'authorization_receipt':'SIM'}
            def send(self,row,auth):
                self.calls+=1
                return {'status':'not_sent','definitely_not_accepted':True} if self.calls==1 else {'status':'sent','message_id':'SIM-OK'}
        with self.s.transaction() as s:core.enqueue(s,'a@example.invalid','提醒','正文','retry')
        host=Host();self.assertEqual(dispatch(self.s,host)['明确失败'],1)
        self.assertEqual(dispatch(self.s,host)['发送成功'],1);dispatch(self.s,host);self.assertEqual(host.calls,2)
    def test_changed_attachment_is_not_dispatched(self):
        class Host:
            EXECUTION_MODE='simulation'
            def authorize(self,row):return {'directory_verified':True,'allowed':True,'authorization_receipt':'SIM'}
            def send(self,row,auth):raise AssertionError('不应调用send')
        import hashlib
        p=self.s.root/'attachment.xlsx';p.write_bytes(b'original')
        with self.s.transaction() as s:core.enqueue(s,'a@example.invalid','报告','正文','attachment',attachments=[{'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}])
        p.write_bytes(b'changed');dispatch(self.s,Host())
        with self.s.transaction() as s:self.assertEqual(s.data['outbox'][0]['状态'],'已失效')
    def test_unavailable_attachment_blocks_cursor_then_resumes(self):
        from docx import Document
        class Host:
            path=None
            def fetch_mail(self,cursor):return {'messages':[{'id':'attachment','sender':'d1@example.invalid','sent_at':now(),'body':'','attachments':[{'name':'evidence.docx','local_path':self.path}]}],'next_cursor':'next'}
            def plan_feedback(self,mail,context):return []
        with self.s.transaction() as s:self.contacts(s)
        host=Host();self.assertEqual(sync(self.s,host)['status'],'incomplete')
        with self.s.transaction() as s:self.assertIsNone(s.get_config('mail_cursor'));self.assertEqual(s.data['mail'],[])
        p=self.s.root/'evidence.docx';doc=Document();doc.add_paragraph('附件内容');doc.save(p);host.path=str(p)
        self.assertEqual(sync(self.s,host)['status'],'ok')
        with self.s.transaction() as s:self.assertEqual(s.get_config('mail_cursor'),'next');self.assertEqual(s.data['pending'][0]['状态'],'自动恢复')

    def test_v1_legacy_ai_is_not_reported_as_current(self):
        with self.s.transaction() as s:
            t=self.task(s);t['当前AI风险等级']='高';t['可能拖延原因']='前版保留的结论'
        with self.s.transaction() as s:
            task=s.data['active'][0]
            self.assertEqual(task['当前AI风险等级'],'未研判')
            self.assertEqual(task['可能拖延原因'],'前版保留的结论')
            self.assertEqual(core.stats(s.data['active'])['高风险任务'],0)

    def test_stale_queued_reports_cannot_bypass_freshness_check(self):
        class Host:
            EXECUTION_MODE='simulation'
            def authorize(self,row): raise AssertionError('陈旧报告不应请求发送授权')
            def send(self,row,auth): raise AssertionError('陈旧报告不应发送')
        with self.s.transaction() as s:self.contacts(s);t=self.task(s);generate(s);queue_reports(s)
        with self.s.transaction() as s:core.update(s,t['内部任务唯一ID'],{'最新办理进展':'报告生成后的新进展'})
        result=dispatch(self.s,Host());self.assertEqual(result['发送成功'],0)
        with self.s.transaction() as s:self.assertTrue(all(r['状态']=='已失效' for r in s.data['outbox']))

if __name__=='__main__':unittest.main()
