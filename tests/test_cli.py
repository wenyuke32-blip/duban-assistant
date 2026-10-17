import json, subprocess, sys, tempfile, unittest
from pathlib import Path
from duban.store import Store, now
from duban import core
from duban.__main__ import sync, extract
from test_business import FakeMail

class CLI(unittest.TestCase):
    def test_full_offline_and_restart(self):
        with tempfile.TemporaryDirectory() as tmp:
            def call(*args):
                p=subprocess.run([sys.executable,'-m','duban',*args,'--workspace',tmp],capture_output=True,text=True)
                self.assertEqual(p.returncode,0,p.stdout+p.stderr)
                return json.loads(p.stdout)
            call('init'); call('add','--file','examples/新增任务.json')
            first=call('query')['结果']['tasks'][0]['内部任务唯一ID']
            result=call('full','--local')
            self.assertEqual(result['结果']['邮箱同步']['status'],'blocked')
            self.assertEqual(len(result['结果']['报告']),10)
            self.assertEqual(result['执行摘要']['本次发送成功数量'],0)
            call('init'); self.assertEqual(call('query')['结果']['tasks'][0]['内部任务唯一ID'],first)
    def test_excel_import_and_docx_pdf_extract(self):
        from openpyxl import Workbook
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp); wb=Workbook(); ws=wb.active
            ws.append(['事项名称','责任部门','截止日期']); ws.append(['审批材料','综合部','2099-01-01']); wb.save(p/'import.xlsx')
            r=subprocess.run([sys.executable,'-m','duban','import','--workspace',str(p/'data'),'--file',str(p/'import.xlsx')],capture_output=True,text=True)
            self.assertEqual(r.returncode,0,r.stdout); self.assertEqual(json.loads(r.stdout)['结果']['新增'],1)
            d=Document(); d.add_paragraph('实际附件测试文字'); d.save(p/'attachment.docx')
            self.assertIn('实际附件测试文字',extract(p/'attachment.docx'))
            self.assertEqual(extract(p/'import.xlsx')['Sheet'][1][0],'审批材料')
    def test_mail_cursor_checkpoint_and_retry(self):
        class Host:
            def __init__(self): self.fail=True
            def fetch_mail(self,cursor):
                return {'messages':[{'id':'cursor-test','sender':'supervisor@example.invalid','sent_at':now(),'body':'专线项目受阻'}],'next_cursor':'next'}
            def plan_feedback(self,mail,context):
                if self.fail: raise RuntimeError()
                t=context['tasks'][0]; return [{'action':'update','task_id':t['内部任务唯一ID'],'candidate_ids':[t['内部任务唯一ID']],'evidence':['专线项目受阻'],'anchors':['专线项目'],'fields':{'办理状态':'受阻'}}]
        with tempfile.TemporaryDirectory() as tmp:
            s=Store(tmp)
            with s.transaction():
                core.contacts(s,{'部门名称':'政企部','主管邮箱':'supervisor@example.invalid'})
                core.create(s,{'任务名称':'专线项目','主责部门':'政企部','原定截止日期':'2099-01-01'})
            host=Host(); self.assertEqual(sync(s,host)['status'],'incomplete')
            with s.transaction(): self.assertIsNone(s.get_config('mail_cursor'))
            host.fail=False; self.assertEqual(sync(s,host)['status'],'ok')
            with s.transaction(): self.assertEqual(s.get_config('mail_cursor'),'next'); self.assertEqual(len(s.data['mail']),1)
            self.assertEqual(sync(s,host)['处理'][0]['重复邮件'],1)
    def test_contact_change_invalidates_old_queue(self):
        with tempfile.TemporaryDirectory() as tmp:
            s=Store(tmp)
            with s.transaction():
                core.enqueue(s,'old@example.invalid','测试','测试','test-key')
                core.contacts(s,{'部门名称':'网络部','主管邮箱':'new@example.invalid'})
                self.assertEqual(s.data['outbox'][0]['状态'],'已失效')

if __name__=='__main__': unittest.main()
