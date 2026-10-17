from __future__ import annotations
import copy, fcntl, json, os, shutil, tempfile, uuid
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

TZ = timezone(timedelta(hours=8))
def now(): return datetime.now(TZ).isoformat(timespec='seconds')
def uid(): return uuid.uuid4().hex

def stamp(value):
    d = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
    return d.replace(tzinfo=TZ) if d.tzinfo is None else d.astimezone(TZ)

FIELDS = ('内部任务唯一ID 任务名称 任务内容 主责部门 责任人 部门负责人 联系邮箱 协办部门 下达日期 原定截止日期 预计完成日期 '
          '办理状态 完成百分比 最新办理进展 存在问题 延期原因 实际完成日期 最近更新时间 最近反馈来源 '
          '当前AI风险等级 可能拖延原因 AI分析依据 AI整改建议 需要协调事项 最新AI研判时间 '
          '最近邮件通知时间 最近通知等级 重新开启次数 累计延期次数 更新时间 创建时间 数据来源').split()
TABLES = {
 'active': ('在办督办任务', FIELDS), 'closed': ('已办结任务', FIELDS),
 'history': ('最近72小时历史记录', ['ID','时间','任务ID','类型','内容']),
 'contacts': ('部门联系人', ['部门名称','主管姓名','主管邮箱','管理员姓名','管理员邮箱']),
 'mail': ('邮件同步及去重记录', ['ID','时间','邮件ID','状态','内容']),
 'pending': ('待确认事项', ['ID','时间','类型','状态','内容']),
 'config': ('系统配置与运行状态', ['键','值']),
 'outbox': ('待发送及发送记录', ['ID','时间','任务ID','接收人','等级','去重键','状态','尝试次数','内容']),
 'runs': ('执行记录', ['ID','时间','内容']),
}
DEPTS = ['市场部','政企部','网络部','综合部']

def safe_cell(v):
    if isinstance(v,(dict,list)): v=json.dumps(v,ensure_ascii=False)
    if isinstance(v,str):
        if len(v)>32000: raise ValueError('单元格内容过长，请拆分输入')
    return v

def style(ws):
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:
        c.font=Font(color='FFFFFF',bold=True); c.fill=PatternFill('solid',fgColor='185A7D')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(48,max(18,len(str(col[0].value))*2+3))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.data_type=='f': c.data_type='s'  # 强制文本，不改变业务原值
            c.alignment=Alignment(vertical='top',wrap_text=True)

class Store:
    def __init__(self, root):
        self.root=Path(root).resolve(); self.root.mkdir(parents=True,exist_ok=True)
        self.path=self.root/'督办管理总台账.xlsx'
        self.backup=self.root/'.recovery'/'督办管理总台账.last-good.xlsx'
        self.data={}
    def read(self,path=None):
        wb=load_workbook(path or self.path, data_only=False)
        data={}
        for k,(name,headers) in TABLES.items():
            if name not in wb.sheetnames: raise ValueError('台账缺少工作表：'+name)
            rows=list(wb[name].values)
            if not rows or list(rows[0])!=headers: raise ValueError('台账表头不兼容：'+name)
            data[k]=[]
            for values in rows[1:]:
                if all(v is None for v in values): continue
                row={h:('' if v is None else v) for h,v in zip(headers,values)}
                for h in ['内容','值']:
                    if h in row and isinstance(row[h],str):
                        try: row[h]=json.loads(row[h])
                        except (ValueError,TypeError): pass
                data[k].append(row)
        wb.close()
        return data
    def save(self):
        wb=Workbook(); wb.remove(wb.active)
        for k,(name,headers) in TABLES.items():
            ws=wb.create_sheet(name); ws.append(headers)
            for row in self.data[k]: ws.append([safe_cell(row.get(h,'')) for h in headers])
            style(ws)
        fd,tmp=tempfile.mkstemp(prefix='.ledger-',suffix='.xlsx',dir=self.root); os.close(fd)
        try:
            wb.save(tmp); wb.close(); self.read(tmp)
            with open(tmp,'rb') as f: os.fsync(f.fileno())
            if self.path.exists():
                self.read()  # 不把已损坏台账覆盖到有效备份
                self.backup.parent.mkdir(exist_ok=True)
                btmp=self.backup.with_suffix('.tmp')
                shutil.copy2(self.path,btmp); os.replace(btmp,self.backup)
            os.replace(tmp,self.path)
            fd=os.open(self.root,os.O_DIRECTORY); os.fsync(fd); os.close(fd)
        finally:
            if os.path.exists(tmp): os.unlink(tmp)
    @contextmanager
    def transaction(self):
        with open(self.root/'.ledger.lock','a') as lock:
            fcntl.flock(lock,fcntl.LOCK_EX)
            if self.path.exists(): self.data=self.read()
            else:
                self.data={k:[] for k in TABLES}
                self.data['contacts']=[{'部门名称':d} for d in DEPTS]
                self.data['config']=[{'键':'schema_version','值':1}]
            cutoff=datetime.now(TZ)-timedelta(hours=72)
            self.data['history']=[r for r in self.data['history'] if stamp(r['时间'])>=cutoff]
            try:
                yield self
                cutoff=datetime.now(TZ)-timedelta(hours=72)
                self.data['history']=[r for r in self.data['history'] if stamp(r['时间'])>=cutoff]
                self.save()
            finally: fcntl.flock(lock,fcntl.LOCK_UN)
    def pending(self,kind,body):
        row={'ID':uid(),'时间':now(),'类型':kind,'状态':'待确认','内容':body}
        self.data['pending'].append(row); return row['ID']
    def history(self,task,kind,body):
        self.data['history'].append({'ID':uid(),'时间':now(),'任务ID':task,'类型':kind,'内容':copy.deepcopy(body)})
    def get_config(self,key,default=None):
        return next((r['值'] for r in self.data['config'] if r['键']==key),default)
    def set_config(self,key,value):
        self.data['config']=[r for r in self.data['config'] if r['键']!=key]
        self.data['config'].append({'键':key,'值':copy.deepcopy(value)})
