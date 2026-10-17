from __future__ import annotations
import argparse, copy, fcntl, hashlib, json, sys
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from .store import Store, now, uid, stamp, TZ
from . import core
from .adapters import Adapter, dispatch, reconcile
from .reports import generate, queue_reports

ALIASES={'初始化督办系统':'init','导入督办台账':'import','新增督办任务':'add','更新督办进度':'update','同步督办邮件':'sync','扫描督办风险':'scan','执行督办预警':'warn','生成部门周报':'reports','生成综合报告':'reports','发送部门周报':'send-reports','查询督办任务':'query','查看待确认事项':'pending','执行完整督办流程':'full'}

def read_json(path): return json.loads(Path(path).read_text(encoding='utf-8'))

def action(s,op):
    name=op.get('action'); fields=op.get('fields',{})
    if name=='create': return core.create(s,fields)
    if name=='update': return core.update(s,op['task_id'],fields)
    if name=='delete': return core.delete(s,op['task_id'])
    if name=='contacts': return core.contacts(s,fields)
    if name=='query':
        filters=op.get('filters',{})
        if set(filters)-set(core.FIELDS): raise ValueError('查询字段无效')
        return [t for t in core.all_tasks(s) if all(t.get(k)==v for k,v in filters.items())]
    if name=='analysis': return core.apply_analysis(s,op['task_id'],op['result'])
    raise ValueError('未支持的结构化操作：'+str(name))

def batch(s,ops):
    results=[]
    if not isinstance(ops,list): raise ValueError('操作计划应为数组')
    for i,op in enumerate(ops):
        previous=copy.deepcopy(s.data)
        try: results.append({'index':i,'ok':True,'result':action(s,op)})
        except Exception as e:
            s.data=previous
            pid=s.pending('管理员操作待确认',{'操作':op,'原因':str(e)})
            results.append({'index':i,'ok':False,'pending_id':pid,'error':str(e)})
    return results

def sync(store,adapter,max_pages=100):
    if not adapter: return {'status':'blocked','读取邮件数量':0,'原因':'缺少已验证的cmcc-mail宿主适配器；未读取邮箱'}
    summary={'status':'ok','读取邮件数量':0,'处理':[],'页数':0,'执行模式':adapter.execution_mode if isinstance(adapter,Adapter) else getattr(adapter,'EXECUTION_MODE','unverified')}
    for _ in range(max_pages):
        with store.transaction() as s: cursor=s.get_config('mail_cursor')
        try:
            page=adapter.fetch_mail(cursor)
            if not isinstance(page,dict) or not isinstance(page.get('messages'),list): raise ValueError('邮件页格式无效')
        except Exception as e:
            with store.transaction() as s: s.pending('邮件拉取失败',{'错误类型':type(e).__name__})
            summary.update({'status':'incomplete','原因':'邮件接口失败，本轮暂停预警'}); return summary
        summary['页数']+=1; summary['读取邮件数量']+=len(page['messages']); complete=True
        def order(mail):
            try: return stamp(mail['sent_at'])
            except Exception: return datetime.min.replace(tzinfo=TZ)
        for original in sorted(page['messages'],key=order):
            mail=copy.deepcopy(original); mid=mail.get('id','未知') if isinstance(mail,dict) else '未知'
            try:
                with store.transaction() as s:
                    if any(r['邮件ID']==mid for r in s.data['mail']):
                        summary['处理'].append({'邮件ID':mid,'重复邮件':1}); continue
                    if not isinstance(mail,dict): raise ValueError('邮件信封无效')
                    sources={'body':str(mail.get('body',''))}
                    digests={}
                    for index,a in enumerate(mail.get('attachments',[])):
                        if not a.get('local_path'): raise ValueError('附件未由合法宿主工具下载')
                        path=Path(a['local_path'])
                        if not path.is_file() or path.stat().st_size>16*1024*1024: raise ValueError('附件不存在或超过16MiB')
                        source_id='attachment:'+str(index)
                        content=extract(path)
                        sources[source_id]=content if isinstance(content,str) else json.dumps(content,ensure_ascii=False)
                        digests[source_id]=hashlib.sha256(path.read_bytes()).hexdigest()
                    mail['evidence_sources']=sources; mail['attachment_digests']=digests
                    context={'tasks':copy.deepcopy(core.all_tasks(s)),'contacts':copy.deepcopy(s.data['contacts'])}
                    plan=adapter.plan_feedback(copy.deepcopy(mail),context)
                    processed=core.process_mail(s,mail,plan)
                    for pending in s.data['pending']:
                        if pending['类型']=='邮件同步失败' and pending['状态']=='待确认' and pending['内容'].get('邮件ID')==mid: pending['状态']='自动恢复'
                summary['处理'].append({'邮件ID':mid,**processed})
            except Exception as e:
                complete=False
                with store.transaction() as s: s.pending('邮件同步失败',{'邮件ID':mid,'错误类型':type(e).__name__})
                summary['处理'].append({'邮件ID':mid,'失败':type(e).__name__})
        if not complete:
            summary.update({'status':'incomplete','原因':'存在失败邮件或未读附件，游标未推进，本轮暂停预警'}); return summary
        next_cursor=page.get('next_cursor',cursor)
        if page.get('has_more') is True and next_cursor==cursor:
            summary.update({'status':'incomplete','原因':'分页游标未推进，本轮暂停预警'}); return summary
        with store.transaction() as s:
            s.set_config('mail_cursor',next_cursor)
            for pending in s.data['pending']:
                if pending['类型']=='邮件拉取失败' and pending['状态']=='待确认': pending['状态']='自动恢复'
        if page.get('has_more') is not True: return summary
    summary.update({'status':'incomplete','原因':'超过最大分页数，本轮暂停预警；再次同步继续'}); return summary

def extract(path):
    p=Path(path); suffix=p.suffix.lower()
    if suffix=='.docx':
        from docx import Document
        doc=Document(p)
        return '\n'.join([x.text for x in doc.paragraphs]+['\t'.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
    if suffix=='.pdf':
        from pypdf import PdfReader
        text='\n'.join(page.extract_text() or '' for page in PdfReader(p).pages)
        if not text.strip(): raise ValueError('PDF无可提取文本，需要合法OCR能力或人工转录')
        return text
    if suffix=='.xlsx':
        wb=load_workbook(p,data_only=True); result={ws.title:[[str(v) if v is not None else '' for v in row] for row in ws.values] for ws in wb}; wb.close(); return result
    raise ValueError('仅支持xlsx/docx/文本型pdf；旧xls/doc需先合法转换')

def run(args):
    store=Store(args.workspace); adapter=Adapter(args.adapter) if args.adapter else None
    cmd=ALIASES.get(args.command,args.command)
    if cmd=='recover': return {'命令':'recover','结果':store.recover()}
    if cmd=='doctor': return {'命令':'doctor','宿主适配器':adapter.capabilities() if adapter else '未配置；真实AI和邮箱未接入','真实平台验收':'未执行'}
    with store.transaction() as s:
        baseline={k:len(s.data[k]) for k in ['active','closed','pending','outbox','mail']}
        old_ids={r['ID'] for r in s.data['history']}
        old_sent={r['ID'] for r in s.data['outbox'] if r['状态']=='已发送'}
    result={}
    if cmd=='sync': result=sync(store,adapter)
    elif cmd=='dispatch':
        result=dispatch(store,adapter) if adapter else {'status':'blocked','原因':'缺少宿主邮箱适配器，保留待发送队列'}
    elif cmd=='reconcile':
        result=reconcile(store,adapter) if adapter else {'status':'blocked','原因':'需要邮箱已发送查询接口'}
    elif cmd=='extract': result=extract(args.file)
    elif cmd in ['warn','full']:
        synced=sync(store,adapter)
        result={'邮箱同步':synced}
        with store.transaction() as s:
            result['风险扫描']=core.scan(s,adapter)
            if synced['status']=='ok' or args.local:
                result['入队预警']=core.warnings(s)
                if args.local: result['本地运行说明']='管理员明确使用--local；未保证已同步邮箱，未真实发信'
            else: result['预警']='暂停，先完成真实邮箱同步；仅离线验收可显式使用--local'
            if cmd=='full':
                result['报告']=generate(s)
                result['报告待发送']=queue_reports(s)
        if adapter and synced['status']=='ok': result['发送']=dispatch(store,adapter)
    else:
        with store.transaction() as s:
            if cmd=='init': result={'台账':str(s.path),'已有任务数':len(core.all_tasks(s)),'联系人':'请用contacts配置四部门主管和同一管理员邮箱','邮箱':'待验证' if not adapter else '配置了适配器，仍需真实联调'}
            elif cmd=='add': result=core.create(s,read_json(args.file))
            elif cmd=='update': result=core.update(s,args.task_id,read_json(args.file))
            elif cmd=='delete': result=core.delete(s,args.task_id)
            elif cmd=='reopen': result=core.update(s,args.task_id,{'办理状态':'办理中'})
            elif cmd=='close': result=core.update(s,args.task_id,{'办理状态':'已完成'})
            elif cmd=='configure':
                config=read_json(args.file)
                if set(config)-{'正式下达邮箱'}: raise ValueError('仅允许配置正式下达邮箱；发送授权不能通过本配置设置')
                senders=config.get('正式下达邮箱',[])
                if not isinstance(senders,list) or any(not core.EMAIL.fullmatch(x) for x in senders): raise ValueError('下达邮箱列表无效')
                s.set_config('正式下达邮箱',senders); result={'正式下达邮箱':senders}
            elif cmd=='contacts': result=core.contacts(s,read_json(args.file))
            elif cmd=='batch': result=batch(s,read_json(args.file))
            elif cmd=='nl':
                if not adapter: result={'status':'blocked','原因':'需宿主AI按prompts/admin.md输出结构化计划后使用batch，或配置plan_admin适配器'}
                else:
                    plan=adapter.plan_admin(args.text,{'tasks':copy.deepcopy(core.all_tasks(s)),'contacts':copy.deepcopy(s.data['contacts'])})
                    result=batch(s,plan)
            elif cmd=='context': result={'tasks':core.all_tasks(s),'contacts':s.data['contacts'],'history_72h':s.data['history'],'today':now()[:10]}
            elif cmd=='query':
                rows=core.all_tasks(s)
                if args.department: rows=[t for t in rows if t['主责部门']==args.department]
                if args.risk: rows=[t for t in rows if t['当前AI风险等级'] in ['高','紧急']]
                result={'tasks':rows,'statistics':core.stats(rows),'responsible':core.grouped(rows,'责任人'),'departments':core.grouped(rows,'主责部门')}
            elif cmd=='pending': result=s.data['pending']
            elif cmd=='resolve':
                pending=next((r for r in s.data['pending'] if r['ID']==args.pending_id and r['状态']=='待确认'),None)
                if not pending: raise ValueError('待确认事项不存在或已处理')
                previous=copy.deepcopy(s.data)
                result=batch(s,read_json(args.file))
                if not all(r['ok'] for r in result): s.data=previous; raise ValueError('确认操作未全部成功，原事项仍待确认')
                pending['状态']='已处理'
            elif cmd=='scan': result=core.scan(s,adapter,True)
            elif cmd=='analysis': result=core.apply_analysis(s,args.task_id,read_json(args.file))
            elif cmd=='reports': result=generate(s)
            elif cmd=='send-reports': result={'入队':queue_reports(s),'下一步':'dispatch执行通讯录核验与平台发送授权'}
            elif cmd=='outbox': result=s.data['outbox']
            elif cmd=='import':
                wb=load_workbook(args.file,data_only=True); ws=wb[args.sheet] if args.sheet else wb.active
                values=list(ws.values); wb.close()
                headers=[str(x or '') for x in values[0]]
                rows=[dict(zip(headers,v)) for v in values[1:] if any(x is not None for x in v)]
                if args.mapping: mapping=read_json(args.mapping)
                else:
                    aliases={'事项名称':'任务名称','督办事项':'任务名称','责任部门':'主责部门','部门':'主责部门','截止日期':'原定截止日期','完成时限':'原定截止日期','进展':'最新办理进展','状态':'办理状态','负责人':'责任人'}
                    mapping={h:(h if h in core.FIELDS else aliases[h]) for h in headers if h in core.FIELDS or h in aliases}
                    if adapter: mapping=adapter.map_columns(headers,rows[:3],core.FIELDS)
                if not isinstance(mapping,dict) or any(v not in core.FIELDS for v in mapping.values()) or len(set(mapping.values()))!=len(mapping): raise ValueError('列映射无效或目标字段重复')
                result=core.import_rows(s,rows,mapping)
            else: raise ValueError('未知命令：'+cmd)
    with store.transaction() as s:
        changes=[r for r in s.data['history'] if r['ID'] not in old_ids]
        if cmd in ['add','update','batch','nl','import','reopen','resolve','sync'] and changes:
            if adapter: core.scan(s,adapter)
        needs=[t['内部任务唯一ID'] for t in s.data['active'] if not core.current_analysis(s,t)]
        summary={'本次新增任务数':sum(r['类型']=='新增' for r in changes),'本次更新任务数':sum(r['类型']=='更新' for r in changes),
                 '本次办结任务数':sum(r['类型']=='办结' for r in changes),'本次重新开启任务数':sum(r['类型']=='重新开启' for r in changes),
                 '本次发送成功数量':sum(r['状态']=='已发送' and r['ID'] not in old_sent and r['内容'].get('发送凭据',{}).get('execution_mode')=='production' for r in s.data['outbox']),
                 '本次模拟发送数量':sum(r['状态']=='已发送' and r['ID'] not in old_sent and r['内容'].get('发送凭据',{}).get('execution_mode')=='simulation' for r in s.data['outbox']),
                 '本次未经环境核实的适配器成功数':sum(r['状态']=='已发送' and r['ID'] not in old_sent and r['内容'].get('发送凭据',{}).get('execution_mode') not in ['production','simulation'] for r in s.data['outbox']),
                 '需宿主AI研判任务ID':needs,
                 '本次新增邮件处理记录':len(s.data['mail'])-baseline['mail'],'本次新增待确认数':len(s.data['pending'])-baseline['pending'],
                 '当前高风险在办任务数':core.stats(s.data['active'])['高风险任务'],
                 '当前待发送数':sum(r['状态'] in ['待授权','失败'] for r in s.data['outbox']),
                 '当前发送结果未知数':sum(r['状态'] in ['发送中','结果未知'] for r in s.data['outbox'])}
        s.data['runs'].append({'ID':uid(),'时间':now(),'内容':{'command':cmd,'summary':summary}})
    return {'命令':cmd,'结果':result,'执行摘要':summary}

def main():
    p=argparse.ArgumentParser(description='督办盯梢小助手V1.1；详见SKILL.md')
    p.add_argument('command'); p.add_argument('--workspace',default='./workspace-data'); p.add_argument('--adapter',help='可信宿主Python模块名，不含凭据')
    p.add_argument('--file'); p.add_argument('--mapping'); p.add_argument('--sheet'); p.add_argument('--task-id'); p.add_argument('--pending-id'); p.add_argument('--department'); p.add_argument('--risk',action='store_true'); p.add_argument('--text'); p.add_argument('--local',action='store_true',help='管理员显式离线运行，不代表真实同步或发送')
    args=p.parse_args()
    try:
        workspace=Path(args.workspace).resolve(); workspace.mkdir(parents=True,exist_ok=True)
        with open(workspace/'.run.lock','a') as lock:
            fcntl.flock(lock,fcntl.LOCK_EX)
            payload=run(args)
        print(json.dumps(payload,ensure_ascii=False,indent=2,default=str))
    except Exception as e:
        # 不输出适配器异常文本/堆栈，防止外部接口把凭据带入日志。
        safe=str(e) if isinstance(e,(ValueError,FileNotFoundError)) else type(e).__name__
        try:
            with Store(args.workspace).transaction() as s:
                s.data['runs'].append({'ID':uid(),'时间':now(),'内容':{'command':args.command,'status':'failed','error_type':type(e).__name__}})
        except Exception: pass
        print(json.dumps({'status':'failed','error':safe,'恢复提示':'原文件写入失败时保留有效台账；检查.recovery后按docs/部署与恢复.md恢复。'},ensure_ascii=False)); sys.exit(1)

if __name__=='__main__': main()
