from __future__ import annotations
import hashlib, json, os, shutil, tempfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt, font_manager
from .store import DEPTS, FIELDS, TZ, now, style, safe_cell
from .core import all_tasks, stats, CLOSED, STATES, RISKS, grouped, enqueue, current_analysis

def chart_font():
    candidates=[f.fname for f in font_manager.fontManager.ttflist if 'Noto Sans CJK' in f.name or 'SimHei' in f.name or 'Microsoft YaHei' in f.name]
    if not candidates: raise RuntimeError('缺少中文图表字体，请安装思源黑体/Noto Sans CJK/微软雅黑')
    return font_manager.FontProperties(fname=candidates[0])

def add_table(ws,headers,rows):
    ws.append(headers)
    for row in rows: ws.append([safe_cell(x) for x in row])
    style(ws)

def build_scope(s,scope,folder):
    tasks=all_tasks(s); own=tasks if scope=='管理员综合' else [t for t in tasks if t['主责部门']==scope]
    assist=[] if scope=='管理员综合' else [t for t in tasks if scope in t.get('协办部门','').split('、') and t['主责部门']!=scope]
    metrics=stats(own); live=[t for t in own if t['办理状态'] not in CLOSED]
    overdue=[t for t in live if t['原定截止日期']<datetime.now(TZ).date().isoformat()]
    risky=[t for t in live if t['当前AI风险等级'] in ['高','紧急'] or t['办理状态']=='受阻']
    analyses={t['内部任务唯一ID']:(current_analysis(s,t) or {}) for t in own}
    status=Counter(t['办理状态'] for t in own)
    risk=Counter(t['当前AI风险等级'] or '未研判' for t in live)
    causes=Counter(analyses[t['内部任务唯一ID']].get('category','尚未研判') for t in live)
    wb=Workbook(); dash=wb.active; dash.title='可视化看板'
    add_table(dash,['指标','数值'],list(metrics.items()))
    for r in range(2,dash.max_row+1):
        if dash.cell(r,1).value in ['完成率','按期完成率']: dash.cell(r,2).number_format='0.00%'
    specs=[('状态分布',[(v,status[v]) for v in STATES]),('风险分布',[(v,risk[v]) for v in RISKS]),('拖延原因',list(causes.items()) or [('无在办任务',0)]),
           ('部门完成情况',[(d,sum(t['办理状态']=='已完成' and t['主责部门']==d for t in own)) for d in (DEPTS if scope=='管理员综合' else [scope])])]
    for index,(title,rows) in enumerate(specs):
        ws=wb.create_sheet(title); add_table(ws,['分类','任务数'],rows)
        chart=PieChart() if title in ['状态分布','风险分布'] else BarChart()
        chart.title=title; chart.add_data(Reference(ws,min_col=2,min_row=1,max_row=ws.max_row),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=ws.max_row)); chart.width=16; chart.height=9
        dash.add_chart(chart,f'{"D" if index%2==0 else "M"}{2+(index//2)*19}')
    for title,subset in [('任务清单',own),('重点任务',risky),('逾期任务',overdue),('协办任务',assist)]:
        add_table(wb.create_sheet(title),FIELDS,[[t[k] for k in FIELDS] for t in subset])
    for title,key in [('责任人统计','责任人'),('部门统计','主责部门')]:
        gs=grouped(own,key); headers=[key]+list(metrics)
        add_table(wb.create_sheet(title),headers,[[name]+list(val.values()) for name,val in gs.items()])
    notes=wb.create_sheet('统计口径')
    add_table(notes,['项目','说明'],[
        ['生成时间',now()],['完成率','已完成/有效任务；有效任务排除已取消。空分母显示空值。'],
        ['按期完成率','实际完成日期不晚于原定截止日期的已完成任务/已完成任务'],
        ['部门归属','每项任务仅计入主责部门；协办任务单列不计入本部门总数'],
        ['逾期','当前未完成且未取消、今天晚于原定截止日期；按Asia/Shanghai自然日'],
        ['数据边界','统计来自当前台账快照；无历史趋势。不将AI推测当作已核实原因。']])
    xlsx=folder/(scope+'督办报告.xlsx'); wb.save(xlsx); wb.close(); load_workbook(xlsx).close()
    font=chart_font(); img=folder/(scope+'.png')
    fig,axes=plt.subplots(1,2,figsize=(11,3.7)); fig.set_facecolor('#f4f8fb')
    for ax,(title,rows) in zip(axes,specs[:2]):
        ax.bar([x[0] for x in rows],[x[1] for x in rows],color='#277E9E')
        ax.set_title(title,fontproperties=font); ax.set_xticks(range(len(rows)))
        ax.set_xticklabels([x[0] for x in rows],fontproperties=font,fontsize=9)
        ax.yaxis.get_major_locator().set_params(integer=True)
    fig.tight_layout(); fig.savefig(img,dpi=150); plt.close(fig)
    doc=Document(); normal=doc.styles['Normal']; normal.font.name='Microsoft YaHei'; normal.font.size=Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
    doc.add_heading(scope+'督办工作分析报告',0); doc.add_paragraph('生成时间：'+now())
    modes={s.get_config('analysis_meta:'+t['内部任务唯一ID'],{}).get('execution_mode','unverified') for t in own if current_analysis(s,t)}
    if modes and modes!={'production'}: doc.add_paragraph('研判来源提示：包含模拟或未经真实宿主调用核验的结构化研判结果，不代表已验证真实大模型分析。')
    doc.add_heading('一、总体督办情况',1)
    ratio=lambda x: '无可计算数据' if x is None else f'{x:.2%}'
    doc.add_paragraph(f"当前共有任务{metrics['任务总数']}项，已完成{metrics['已完成数']}项，在办{metrics['未完成数']}项，已取消{metrics['已取消数']}项。完成率{ratio(metrics['完成率'])}，按期完成率{ratio(metrics['按期完成率'])}。当前逾期{metrics['逾期数量']}项，高风险{metrics['高风险任务']}项，受阻{metrics['受阻任务']}项。")
    doc.add_picture(str(img),width=Inches(6.2))
    doc.add_heading('二、主要工作进展',1)
    if not own: doc.add_paragraph('本报告范围内暂无任务记录。')
    for t in own: doc.add_paragraph(f"【{t['办理状态']}】{t['任务名称']}（责任人：{t['责任人'] or '待配置'}；原截止日期：{t['原定截止日期']}）：{t['最新办理进展'] or '尚无进展反馈'}。")
    doc.add_heading('三、重点风险及拖延原因研判',1)
    if not risky: doc.add_paragraph('台账暂无标记为高风险、紧急或受阻的在办任务。该结果不代表已完成所有任务的AI研判。')
    for t in risky:
        doc.add_heading(t['任务名称'],2)
        doc.add_paragraph('最新进展：'+(t['最新办理进展'] or '未反馈')+'；已报告问题：'+(t['存在问题'] or '未反馈'))
        a=analyses[t['内部任务唯一ID']]
        if not a: doc.add_paragraph('尚未取得宿主AI的有效研判结果，请先执行风险扫描。'); continue
        for label,key in [('记录事实','facts'),('待核实可能原因','hypotheses'),('可核对依据','evidence'),('具体措施','actions'),('需协调事项','coordination'),('下一阶段跟踪','next_focus')]:
            doc.add_paragraph(label+'：'+('；'.join(a.get(key,[])) or '无已记录内容'))
    doc.add_heading('四、协调问题、整改措施及下阶段重点',1)
    for t in live:
        a=analyses[t['内部任务唯一ID']]
        doc.add_paragraph(f"{t['任务名称']}：协调事项：{('；'.join(a.get('coordination',[]))) or '尚未有效研判'}；整改建议：{('；'.join(a.get('actions',[]))) or '尚未有效研判'}；下阶段重点：{'；'.join(a.get('next_focus',[])) or '核实最新进展与实际阻塞事项'}。")
    if not live: doc.add_paragraph('当前无在办任务。')
    narrative=s.get_config('narrative:'+scope,{})
    if narrative.get('fingerprint')==business_report_fingerprint(s):
        doc.add_heading('五、MobileWork智能复盘',1)
        if narrative.get('execution_mode')!='production': doc.add_paragraph('以下为模拟或未经宿主实际调用核验的复盘内容，不能作为真实模型验收证据。')
        for section in narrative['sections']:
            doc.add_heading(section['heading'],2);doc.add_paragraph(section['text'])
            if section.get('evidence'):doc.add_paragraph('依据：'+'；'.join(e['quote'] for e in section['evidence']))
    doc.add_heading('六、协办及统计明细',1)
    if scope=='管理员综合':
        table=doc.add_table(rows=1,cols=4); table.style='Light Shading Accent 1'
        for c,v in zip(table.rows[0].cells,['部门','任务总数','已完成','逾期']): c.text=v
        for d in DEPTS:
            m=stats([t for t in own if t['主责部门']==d])
            for c,v in zip(table.add_row().cells,[d,m['任务总数'],m['已完成数'],m['逾期数量']]): c.text=str(v)
    for t in assist: doc.add_paragraph('协办任务：'+t['任务名称']+'；主责部门：'+t['主责部门']+'；协办内容以任务要求为准：'+t['任务内容'])
    doc.add_paragraph('统计口径：任务只计入主责部门；完成率分母排除已取消；按期完成率分母为已完成。详细任务、责任人统计和拖延原因分类见配套Excel。AI可能原因尚待核实。')
    docx=folder/(scope+'督办报告.docx'); doc.save(docx); Document(docx)
    img.unlink()
    return [xlsx,docx]

def business_report_fingerprint(s):
    volatile={'更新时间','最近更新时间','最新AI研判时间','最近邮件通知时间','最近通知等级','最近反馈来源'}
    tasks=[{k:t[k] for k in FIELDS if k not in volatile} for t in sorted(all_tasks(s),key=lambda t:t['内部任务唯一ID'])]
    analyses={t['内部任务唯一ID']:current_analysis(s,t) for t in all_tasks(s)}
    payload={'date':datetime.now(TZ).date().isoformat(),'version':'1.2.0','tasks':tasks,'analyses':analyses,'analysis_modes':{t['内部任务唯一ID']:s.get_config('analysis_meta:'+t['内部任务唯一ID'],{}).get('execution_mode') for t in all_tasks(s)}}
    return hashlib.sha256(json.dumps(payload,ensure_ascii=False,sort_keys=True).encode()).hexdigest()

def report_context(s,scope):
    if scope not in DEPTS+['管理员综合']: raise ValueError('报告范围无效')
    tasks=all_tasks(s)
    own=tasks if scope=='管理员综合' else [t for t in tasks if t['主责部门']==scope]
    assist=[] if scope=='管理员综合' else [t for t in tasks if scope in t.get('协办部门','').split('、') and t['主责部门']!=scope]
    return {'scope':scope,'expected_fingerprint':business_report_fingerprint(s),'statistics':stats(own),
            'tasks':own,'assisting_tasks':assist,'analyses':{t['内部任务唯一ID']:current_analysis(s,t) for t in own},
            '要求':'由MobileWork自身大模型基于这些真实统计与原文撰写带引用的进展、风险、协调、整改、下阶段复盘。数字只取statistics，不自由估算。'}

def save_narrative(s,scope,sections,expected_fingerprint,execution_mode='unverified'):
    context=report_context(s,scope)
    if expected_fingerprint!=context['expected_fingerprint']: raise ValueError('报告上下文已变化，需要重新生成AI复盘')
    if not isinstance(sections,list) or not sections: raise ValueError('报告复盘至少包含一个段落')
    allowed={t['内部任务唯一ID']:t for t in context['tasks']+context['assisting_tasks']}
    other_names=[t['任务名称'] for t in all_tasks(s) if t['内部任务唯一ID'] not in allowed]
    headings=['总体情况','主要进展','风险分析','协调事项','整改建议','下阶段重点']
    for section in sections:
        if not isinstance(section,dict) or section.get('heading') not in headings or not isinstance(section.get('text'),str) or not section['text'].strip() or len(section['text'])>4000: raise ValueError('复盘段落结构无效')
        ids=section.get('task_ids',[]);evidence=section.get('evidence',[])
        if not isinstance(ids,list) or any(tid not in allowed for tid in ids): raise ValueError('复盘引用超出报告范围')
        if allowed and (not ids or not evidence): raise ValueError('有任务时复盘须给出任务ID及原文证据')
        if not isinstance(evidence,list): raise ValueError('证据结构无效')
        for e in evidence:
            task=allowed.get(e.get('task_id'))
            if not task or task['内部任务唯一ID'] not in ids or not isinstance(e.get('quote'),str) or not e['quote'].strip(): raise ValueError('复盘原文引用无效')
            values=[str(v) for k,v in task.items() if k not in ['可能拖延原因','AI分析依据','AI整改建议','需要协调事项'] or current_analysis(s,task)]
            if not any(e['quote'] in value for value in values): raise ValueError('复盘证据不能在当前任务中核对')
        if any(name in section['text'] for name in other_names): raise ValueError('部门复盘包含其他部门专属任务')
    record={'fingerprint':expected_fingerprint,'sections':sections,'execution_mode':execution_mode,'generated_at':now()}
    s.set_config('narrative:'+scope,record)

def report_fingerprint(s):
    base=business_report_fingerprint(s)
    narratives={scope:s.get_config('narrative:'+scope,{}).get('sections') for scope in DEPTS+['管理员综合'] if s.get_config('narrative:'+scope,{}).get('fingerprint')==base}
    return hashlib.sha256(json.dumps({'base':base,'narratives':narratives,'narrative_modes':{scope:s.get_config('narrative:'+scope,{}).get('execution_mode') for scope in narratives}},ensure_ascii=False,sort_keys=True).encode()).hexdigest()

def _sync_dir(root):
    fd=os.open(root,os.O_DIRECTORY)
    try: os.fsync(fd)
    finally: os.close(fd)

def _journal_write(root,journal):
    fd,tmp=tempfile.mkstemp(prefix='.publish-',dir=root)
    try:
        with os.fdopen(fd,'w',encoding='utf-8') as f:
            json.dump(journal,f,ensure_ascii=False); f.flush(); os.fsync(f.fileno())
        os.replace(tmp,root/'.report-publication.json'); _sync_dir(root)
    finally:
        if os.path.exists(tmp): os.unlink(tmp)

def _point(root,target):
    link=root/'reports'; tmp=root/'.reports-next'
    if link.exists() and not link.is_symlink(): raise ValueError('reports路径被外部目录占用')
    if tmp.is_symlink(): tmp.unlink()
    if target is None:
        if link.is_symlink(): link.unlink()
    else:
        tmp.symlink_to(target,target_is_directory=True); os.replace(tmp,link)
    _sync_dir(root)

def recover_publication(s,committed=None):
    journal_path=s.root/'.report-publication.json'
    if not journal_path.exists(): return
    journal=json.loads(journal_path.read_text(encoding='utf-8'))
    for key in ['old','new']:
        name=journal.get(key)
        if name is not None and (not isinstance(name,str) or not name.startswith('.reports-') or Path(name).name!=name):
            raise ValueError('报告恢复日志目录不合法')
    if committed is None: committed=s.get_config('reports',{}).get('hashes')==journal['hashes']
    target=journal['new'] if committed else journal['old']
    if target and not (s.root/target).is_dir(): raise ValueError('报告恢复所需版本缺失，请人工核对')
    _point(s.root,target)
    obsolete=journal['old'] if committed else journal['new']
    if obsolete and obsolete!=target and (s.root/obsolete).exists(): shutil.rmtree(s.root/obsolete,ignore_errors=False)
    journal_path.unlink(); _sync_dir(s.root)

def _manifest_valid(s,manifest):
    folder=s.root/'reports'
    hashes=manifest.get('hashes',{})
    return (len(hashes)==10 and folder.is_dir() and {p.name for p in folder.iterdir()}==set(hashes)
            and all((folder/name).is_file() and hashlib.sha256((folder/name).read_bytes()).hexdigest()==digest for name,digest in hashes.items()))

def generate(s):
    """业务快照未变时复用；发布日志+事务回调保证失败及中断可恢复。"""
    root=s.root; link=root/'reports'; fingerprint=report_fingerprint(s)
    previous=s.get_config('reports',{})
    if previous.get('fingerprint')==fingerprint and _manifest_valid(s,previous):
        return [str(link/name) for name in previous['hashes']]
    if link.exists() and not link.is_symlink(): raise ValueError('reports必须由程序管理；请先迁移旧的手工目录')
    old=link.resolve() if link.is_symlink() else None
    if old and (old.parent!=root or not old.name.startswith('.reports-')): raise ValueError('报告链接指向非程序管理目录')
    stage=Path(tempfile.mkdtemp(prefix='.reports-',dir=root)); published=False
    try:
        files=[]
        for scope in DEPTS+['管理员综合']: files.extend(build_scope(s,scope,stage))
        manifest={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
        for p in files:
            with p.open('rb') as f: os.fsync(f.fileno())
        _sync_dir(stage)
        journal={'old':old.name if old else None,'new':stage.name,'hashes':manifest}
        _journal_write(root,journal)
        # 即使切换或台账save失败，事务退出时也能恢复。
        s._rollback_hooks.append(lambda:recover_publication(s,False))
        s._commit_hooks.append(lambda:recover_publication(s,True))
        _point(root,stage.name); published=True
        s.set_config('reports',{'生成时间':now(),'hashes':manifest,'fingerprint':fingerprint,'version_dir':stage.name})
        return [str(link/p.name) for p in files]
    except BaseException:
        if not (root/'.report-publication.json').exists() and not published: shutil.rmtree(stage,ignore_errors=True)
        raise

def queue_reports(s):
    info=s.get_config('reports',{})
    if info.get('fingerprint')!=report_fingerprint(s) or not _manifest_valid(s,info): raise ValueError('报告与当前台账不一致，请先重新生成')
    manifest=info['hashes']; queued=[]
    for scope in DEPTS+['管理员综合']:
        c=next(r for r in s.data['contacts'] if r['部门名称']==(DEPTS[0] if scope=='管理员综合' else scope))
        recipient=c.get('管理员邮箱','') if scope=='管理员综合' else c.get('主管邮箱','')
        if scope=='管理员综合':
            admins={r.get('管理员邮箱') for r in s.data['contacts'] if r.get('管理员邮箱')}
            if len(admins)!=1:
                s.pending('管理员邮箱未唯一配置',{}); continue
            recipient=next(iter(admins))
        attachments=[]
        for ext in ['xlsx','docx']:
            name=scope+'督办报告.'+ext; p=s.root/'reports'/name
            if not p.exists() or name not in manifest: raise ValueError('报告未生成或缺少清单')
            digest=hashlib.sha256(p.read_bytes()).hexdigest()
            if digest!=manifest[name]: raise ValueError('报告已被外部修改，请重新生成')
            attachments.append({'path':str(p),'sha256':digest})
        key='report|'+scope+'|'+recipient+'|'+hashlib.sha256(json.dumps(attachments,sort_keys=True).encode()).hexdigest()
        q=enqueue(s,recipient,scope+'报告','附件为'+scope+'督办报告。统计口径及AI研判边界详见报告。',key,attachments=attachments)
        if q: queued.append(q)
    return queued
